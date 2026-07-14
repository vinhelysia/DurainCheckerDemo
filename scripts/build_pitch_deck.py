# -*- coding: utf-8 -*-
"""Generate DurianTrust UniHackfest pitch deck (PPTX) — Vietnamese for judges."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as RgbColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

INK = RgbColor(0x16, 0x20, 0x1B)
INK_SOFT = RgbColor(0x46, 0x50, 0x4A)
INK_FAINT = RgbColor(0x6E, 0x79, 0x73)
EMERALD = RgbColor(0x0E, 0x7A, 0x52)
GREEN_DEEP = RgbColor(0x12, 0x3B, 0x2C)
GOLD = RgbColor(0x9A, 0x6A, 0x10)
BG = RgbColor(0xF7, 0xF8, 0xF7)
BG_ALT = RgbColor(0xED, 0xF0, 0xEE)
WHITE = RgbColor(0xFF, 0xFF, 0xFF)
RISK_HIGH = RgbColor(0xB3, 0x36, 0x2A)
MUTED = RgbColor(0xD0, 0xD8, 0xD4)
MUTED2 = RgbColor(0x9A, 0xA8, 0xA2)
SOFT_GREEN = RgbColor(0xA8, 0xC4, 0xB8)
CLOSE_SOFT = RgbColor(0xC5, 0xD4, 0xCE)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "DurianTrust_Pitch_Deck.pptx")
HERO = os.path.join(ROOT, "public", "images", "hero-durian.jpg")

# Prefer fonts that render Vietnamese well on Windows
FONT = "Calibri"
FONT_MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height
TOTAL = 11


def set_run(run, text, size=18, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_round_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_text_box(slide, l, t, w, h, text, size=18, bold=False, color=INK,
                 align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color, font=font)
    return box


def multiline_box(slide, l, t, w, h, lines, size=14, bold=False, color=INK,
                  space_before=6, font=FONT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(space_before if i else 0)
        run = p.add_run()
        set_run(run, line, size=size, bold=bold, color=color, font=font)
    return box


def footer(slide, page):
    add_text_box(
        slide, Inches(0.5), Inches(7.1), Inches(9), Inches(0.3),
        "DurianTrust  ·  UniHackfest RWA  ·  Solana Devnet",
        size=11, color=INK_FAINT,
    )
    add_text_box(
        slide, Inches(11.5), Inches(7.1), Inches(1.4), Inches(0.3),
        f"{page} / {TOTAL}", size=11, color=INK_FAINT, align=PP_ALIGN.RIGHT,
    )


def accent_bar(slide):
    add_rect(slide, Inches(0), Inches(0), Inches(0.12), H, EMERALD)


def section_label(slide, text):
    add_text_box(
        slide, Inches(0.55), Inches(0.35), Inches(8), Inches(0.3),
        text.upper(), size=12, bold=True, color=EMERALD,
    )


# ----- 1 Title -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, GREEN_DEEP)
add_rect(s, 0, 0, Inches(0.18), H, GOLD)
if os.path.exists(HERO):
    try:
        s.shapes.add_picture(HERO, Inches(7.2), 0, height=H)
    except Exception as e:
        print("hero fail", e)

add_text_box(s, Inches(0.7), Inches(1.6), Inches(7), Inches(0.4),
             "UNIHACKFEST  ·  HẠNG MỤC RWA", size=14, bold=True, color=GOLD)
add_text_box(s, Inches(0.7), Inches(2.15), Inches(7.5), Inches(1.2),
             "DurianTrust", size=54, bold=True, color=WHITE)
multiline_box(
    s, Inches(0.7), Inches(3.35), Inches(7.2), Inches(1.2),
    [
        "Xác minh lô sầu riêng xuất khẩu trên blockchain —",
        "từ vườn đến hải quan, trong vài giây thay vì vài ngày.",
    ],
    size=18, color=MUTED, space_before=4,
)
add_text_box(
    s, Inches(0.7), Inches(5.1), Inches(7.2), Inches(0.5),
    "Sổ cái Solana  ·  Cổng an toàn Cadimi  ·  Chuyển giao 2 chữ ký  ·  Quét QR",
    size=13, color=GOLD,
)
add_text_box(
    s, Inches(0.7), Inches(6.5), Inches(7), Inches(0.35),
    "Lô demo  DRN-2026-LD-0429   ·   Devnet đang live",
    size=12, color=MUTED2,
)

# ----- 2 Problem -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Vấn đề")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.9),
    "Trái cây hỏng ở cửa khẩu trong khi giấy tờ còn đang truy vết.",
    size=28, bold=True, color=INK,
)
cards = [
    (
        "Siết Cadimi & phẩm màu",
        "Từ 2025, hải quan Trung Quốc (GACC) siết kiểm soát Cadimi và phẩm vàng O trên sầu riêng Việt Nam.",
    ),
    (
        "Chứng từ manh mún",
        "Phiếu lab giấy/PDF, mã vùng trồng dễ tái sử dụng/giả, hồ sơ chuỗi lạnh không xuất trình kịp.",
    ),
    (
        "Mất ngày để truy nguồn",
        "Khi lô bị nghi, truy về vườn có thể mất nhiều ngày — phần còn lại hỏng tại biên giới.",
    ),
]
for i, (title, body) in enumerate(cards):
    x = Inches(0.55 + i * 4.15)
    add_round_rect(s, x, Inches(2.0), Inches(3.95), Inches(3.6), WHITE)
    add_rect(s, x, Inches(2.0), Inches(3.95), Inches(0.1), RISK_HIGH if i == 2 else GOLD)
    add_text_box(s, x + Inches(0.25), Inches(2.35), Inches(3.4), Inches(0.35),
                 f"0{i+1}", size=14, bold=True, color=EMERALD)
    add_text_box(s, x + Inches(0.25), Inches(2.8), Inches(3.4), Inches(0.9),
                 title, size=18, bold=True, color=INK)
    add_text_box(s, x + Inches(0.25), Inches(3.75), Inches(3.4), Inches(1.6),
                 body, size=13, color=INK_SOFT)
add_text_box(
    s, Inches(0.55), Inches(5.9), Inches(12), Inches(0.5),
    "Nguồn: VietnamPlus · MOIT/VNTR · Tuổi Trẻ · SGGP",
    size=12, color=INK_FAINT,
)
footer(s, 2)

# ----- 3 RWA -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Vì sao đây là RWA")
multiline_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(1.0),
    [
        "Một lô sầu riêng là tài sản thật —",
        "mỗi lần chuyển tay là trách nhiệm pháp lý cũng chuyển theo.",
    ],
    size=26, bold=True, color=INK, space_before=2,
)
rows = [
    ("Lô vật lý", "Hàng xuất khẩu thật ngoài đời"),
    ("Batch PDA", "Bản sao số trên Solana (không mint NFT)"),
    ("Quyền sở hữu (custody)", "Ai đang nắm quyền / chịu trách nhiệm lô này"),
    ("Vai trò (roles)", "Ai được ghi dữ liệu: nông hộ / lab / logistics"),
]
for i, (a, b) in enumerate(rows):
    y = Inches(2.5 + i * 0.85)
    add_round_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.75), WHITE)
    add_rect(s, Inches(0.55), y, Inches(0.12), Inches(0.75), EMERALD)
    add_text_box(s, Inches(0.9), y + Inches(0.18), Inches(4.2), Inches(0.4),
                 a, size=17, bold=True, color=GREEN_DEEP)
    add_text_box(s, Inches(5.2), y + Inches(0.18), Inches(7.2), Inches(0.4),
                 b, size=15, color=INK_SOFT)
footer(s, 3)

# ----- 4 Product -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Sản phẩm")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7),
    "Một lớp tin cậy số cho mọi lô xuất khẩu.",
    size=28, bold=True, color=INK,
)
features = [
    ("Sổ cái chỉ ghi thêm",
     "Thu hoạch, lab, đóng gói, xuất — mỗi bước là giao dịch Solana, không sửa lặng lẽ sau."),
    ("Chuyển giao 2 chữ ký",
     "Đề cử + chấp nhận. Không ép nhận. Admin có thể tạm dừng — không tịch thu lô không thuộc mình."),
    ("Cổng Cadimi 0,05 ppm",
     "Kết quả xét nghiệm đối chiếu ngưỡng ngay khi nhập. Quyết định pass / xem xét / giữ — minh bạch, không hộp đen."),
    ("AI hỗ trợ sớm",
     "Mô hình ONNX: rủi ro Cadimi, bệnh, ảnh lá — chỉ sàng lọc; phiếu lab có chữ ký mới là chuẩn."),
    ("QR cho mọi người",
     "Người mua & hải quan quét là thấy timeline, custody, bằng chứng Explorer. Không cần ví để tra cứu."),
    ("Hai chế độ trung thực",
     "Live on-chain khi có mạng; huy hiệu Demo nếu offline. Không giả ownership trong fallback."),
]
for i, (title, body) in enumerate(features):
    col = i % 3
    row = i // 3
    x = Inches(0.55 + col * 4.2)
    y = Inches(1.7 + row * 2.4)
    add_round_rect(s, x, y, Inches(4.0), Inches(2.2), WHITE)
    add_text_box(s, x + Inches(0.25), y + Inches(0.25), Inches(3.5), Inches(0.5),
                 title, size=16, bold=True, color=EMERALD)
    add_text_box(s, x + Inches(0.25), y + Inches(0.85), Inches(3.5), Inches(1.2),
                 body, size=12, color=INK_SOFT)
footer(s, 4)

# ----- 5 Live demo -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, GREEN_DEEP)
add_rect(s, 0, 0, Inches(0.18), H, GOLD)
add_text_box(s, Inches(0.7), Inches(0.6), Inches(12), Inches(0.35),
             "DEMO TRỰC TIẾP", size=14, bold=True, color=GOLD)
add_text_box(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.8),
             "Xem được ngay — không cần ví.", size=34, bold=True, color=WHITE)
steps = [
    ("1", "Mở app → “Scan a demo batch” / Quét lô demo", "#/unit/demo"),
    ("2", "Mở lô flagship", "DRN-2026-LD-0429"),
    ("3", "Timeline + lab so với ngưỡng Cadimi", "0,05 ppm"),
    ("4", "Chuỗi custody + link Solana Explorer", "Handoff thật trên devnet"),
    ("5", "Tuỳ chọn: Accept custody hoặc ghi lab", "#/manage + Phantom Devnet"),
]
for i, (n, a, b) in enumerate(steps):
    y = Inches(2.15 + i * 0.85)
    add_text_box(s, Inches(0.7), y, Inches(0.6), Inches(0.5),
                 n, size=24, bold=True, color=GOLD)
    add_text_box(s, Inches(1.4), y, Inches(7.8), Inches(0.45),
                 a, size=17, color=WHITE)
    add_text_box(s, Inches(9.2), y, Inches(3.6), Inches(0.45),
                 b, size=13, color=SOFT_GREEN)
add_text_box(
    s, Inches(0.7), Inches(6.55), Inches(12), Inches(0.4),
    "Dự phòng mất Wi‑Fi: ảnh chụp cùng lô + huy hiệu “Demo data” (hai chế độ có chủ đích).",
    size=13, color=MUTED2,
)

# ----- 6 Architecture -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Kiến trúc")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
    "Chain là nguồn sự thật. UI đọc chain. AI chỉ hỗ trợ.",
    size=26, bold=True, color=INK,
)
layers = [
    (GREEN_DEEP, WHITE, "Client",
     ["React 19 + Vite", "Phantom (ghi)", "Quét / tạo QR", "Song ngữ VI / EN"]),
    (EMERALD, WHITE, "Solana Devnet",
     ["Chương trình Anchor", "durian_trust", "Batch · Lab · Timeline", "Custody PDA"]),
    (GOLD, WHITE, "AI (ONNX)",
     ["Rủi ro Cadimi", "Rủi ro bệnh", "Phân loại ảnh lá", "Vercel Python"]),
    (INK_SOFT, WHITE, "Fallback",
     ["Dữ liệu demo sẵn", "Sổ localStorage", "Luôn có huy hiệu", "Không giả ownership"]),
]
for i, (bgc, fgc, title, body_lines) in enumerate(layers):
    x = Inches(0.55 + i * 3.2)
    add_round_rect(s, x, Inches(1.8), Inches(3.0), Inches(4.2), bgc)
    add_text_box(s, x + Inches(0.25), Inches(2.15), Inches(2.5), Inches(0.5),
                 title, size=17, bold=True, color=fgc)
    multiline_box(
        s, x + Inches(0.25), Inches(2.85), Inches(2.5), Inches(2.8),
        body_lines, size=13, color=fgc, space_before=6,
    )
footer(s, 6)

# ----- 7 Guarantees -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Cam kết trên chain")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7),
    "Những quy tắc giám khảo có thể thử ngay.",
    size=26, bold=True, color=INK,
)
guarantees = [
    ("Lab chỉ ghi thêm", "Báo cáo mới nhất là chuẩn; lịch sử vẫn giữ."),
    ("Đăng ký không bị kẹt logistics", "Nông hộ/HTX đăng ký lô trước, không cần logistics."),
    ("Custody 2 chữ ký", "Ví sai không accept; không sở hữu thì không transfer."),
    ("Admin không tịch thu", "Authority có thể pause — không gán lại ownership."),
    ("Khoá genesis", "Initialize chỉ bởi genesis authority hardcode."),
    ("Hai chế độ rõ ràng", "Huy hiệu Live vs Demo — không giả ownership."),
]
for i, (t, b) in enumerate(guarantees):
    col = i % 2
    row = i // 2
    x = Inches(0.55 + col * 6.35)
    y = Inches(1.65 + row * 1.55)
    add_round_rect(s, x, y, Inches(6.1), Inches(1.4), WHITE)
    add_rect(s, x, y, Inches(0.12), Inches(1.4), EMERALD)
    add_text_box(s, x + Inches(0.35), y + Inches(0.25), Inches(5.5), Inches(0.4),
                 t, size=16, bold=True, color=INK)
    add_text_box(s, x + Inches(0.35), y + Inches(0.75), Inches(5.5), Inches(0.45),
                 b, size=13, color=INK_SOFT)
footer(s, 7)

# ----- 8 Practicality -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Thực dụng ngoài đời")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.9),
    "Nông dân không cần “chơi crypto”.",
    size=28, bold=True, color=INK,
)
add_round_rect(s, Inches(0.55), Inches(1.85), Inches(6.0), Inches(4.5), WHITE)
add_text_box(s, Inches(0.85), Inches(2.15), Inches(5.4), Inches(0.45),
             "Demo hackathon (hiện tại)", size=17, bold=True, color=GOLD)
multiline_box(
    s, Inches(0.85), Inches(2.8), Inches(5.4), Inches(3.2),
    [
        "▸  Mỗi vai trò dùng Phantom để chứng minh protocol",
        "▸  Chữ ký thật + bằng chứng trên Explorer",
        "▸  Phù hợp để ban giám khảo kiểm tra on-chain",
    ],
    size=14, color=INK_SOFT, space_before=10,
)
add_round_rect(s, Inches(6.8), Inches(1.85), Inches(6.0), Inches(4.5), GREEN_DEEP)
add_text_box(s, Inches(7.1), Inches(2.15), Inches(5.4), Inches(0.45),
             "Triển khai pilot (tiếp theo)", size=17, bold=True, color=GOLD)
multiline_box(
    s, Inches(7.1), Inches(2.8), Inches(5.4), Inches(3.2),
    [
        "▸  HTX / nhà đóng gói / lab giữ khoá được cấp quyền",
        "▸  Nông hộ dùng form / SMS / tablet — nhận mã lô",
        "▸  Doanh nghiệp trả phí gas (SOL = phí hạ tầng)",
        "▸  Hải quan & người mua chỉ quét QR (chỉ đọc)",
    ],
    size=14, color=WHITE, space_before=10,
)
footer(s, 8)

# ----- 9 Honesty -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Bảng trung thực")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
    "Cái gì là thật — cái gì hỗ trợ — cái gì chưa claim.",
    size=26, bold=True, color=INK,
)
table_data = [
    ("Đăng ký lô / lab / timeline", "Thật trên Solana devnet"),
    ("Chuỗi custody (2 chữ ký)", "Thật — chỉ on-chain, không giả fallback"),
    ("Cổng Cadimi 0,05 ppm", "Hằng số quy tắc, kiểm được"),
    ("Mô hình AI (ONNX)", "Hỗ trợ sàng lọc thật"),
    ("QR + bằng chứng Explorer", "Deep link / chữ ký thật"),
    ("Triển khai toàn quốc nông hộ", "Chưa claim — có lộ trình pilot"),
    ("Mainnet / thông quan thật", "Chưa — demo trên devnet"),
]
for i, (a, b) in enumerate(table_data):
    y = Inches(1.5 + i * 0.7)
    bgc = WHITE if i % 2 == 0 else BG_ALT
    add_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.65), bgc)
    add_text_box(s, Inches(0.75), y + Inches(0.15), Inches(6.5), Inches(0.4),
                 a, size=14, bold=True, color=INK)
    add_text_box(s, Inches(7.4), y + Inches(0.15), Inches(5.1), Inches(0.4),
                 b, size=13, color=INK_SOFT)
footer(s, 9)

# ----- 10 Hard facts -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, BG)
accent_bar(s)
section_label(s, "Thông số cứng cho giám khảo")
add_text_box(
    s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
    "Chính xác — không tự bịa ID.",
    size=26, bold=True, color=INK,
)
facts = [
    ("Program ID", "4EZcqRn9LYK5VMuhLC2bNDaUqVBHxc6KCZ6zhFet3Par"),
    ("Mạng", "Solana Devnet"),
    ("Lô flagship", "DRN-2026-LD-0429"),
    ("Ngưỡng Cadimi", "0,05 ppm"),
    ("Genesis authority", "52WpskyDdHaLyAcyTLQrqvLBUh3azKFAe3XmNkYDaFJu"),
    ("Ví demo", "2BARgkoYQPL7ngMerfCh21CpGRepuZdVtMGr7do1ssko"),
    ("Đường dẫn app", "#/   ·   #/unit/demo   ·   #/manage"),
]
for i, (k, v) in enumerate(facts):
    y = Inches(1.5 + i * 0.7)
    add_round_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.62), WHITE)
    add_text_box(s, Inches(0.8), y + Inches(0.12), Inches(3.2), Inches(0.4),
                 k, size=14, bold=True, color=EMERALD)
    add_text_box(s, Inches(4.1), y + Inches(0.12), Inches(8.4), Inches(0.4),
                 v, size=13, color=INK, font=FONT_MONO)
footer(s, 10)

# ----- 11 Close -----
s = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(s, 0, 0, W, H, GREEN_DEEP)
add_rect(s, 0, 0, Inches(0.18), H, GOLD)
add_text_box(s, Inches(0.7), Inches(1.6), Inches(12), Inches(0.4),
             "Câu kết", size=14, bold=True, color=GOLD)
multiline_box(
    s, Inches(0.7), Inches(2.15), Inches(12), Inches(2.0),
    [
        "Bản sao số của lô xuất khẩu.",
        "Trách nhiệm đi cùng quyền custody.",
        "Cổng an toàn có thể kiểm toán.",
        "QR cho mọi người còn lại.",
    ],
    size=26, bold=True, color=WHITE, space_before=4,
)
add_text_box(
    s, Inches(0.7), Inches(4.5), Inches(12), Inches(0.7),
    "Blockchain là lớp tin cậy phía sau — không phải app crypto cho nông hộ.",
    size=17, color=CLOSE_SOFT,
)
add_text_box(
    s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.4),
    "DurianTrust  ·  Quét lô demo DRN-2026-LD-0429  ·  Xin cảm ơn",
    size=15, color=GOLD,
)
add_text_box(s, Inches(0.7), Inches(6.4), Inches(12), Inches(0.35),
             "Hỏi & đáp", size=14, color=MUTED2)

prs.save(OUT)
print("SAVED", OUT)
print("slides", len(prs.slides))
